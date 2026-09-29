"""Builds real sample files for the document tests — genuine PDFs, DOCX and PPTX produced by the
same libraries users' tools produce, not hand-faked bytes."""

from io import BytesIO

from docx import Document as new_docx
from fpdf import FPDF
from PIL import Image
from pptx import Presentation


def make_pdf(
    pages: list[tuple[str | None, str]],
    *,
    user_password: str | None = None,
    owner_password: str | None = None,
) -> bytes:
    """One page per (bookmark title or None, text). Titles become PDF bookmarks (the outline)."""
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for title, text in pages:
        pdf.add_page()
        if title:
            pdf.start_section(title)
        pdf.multi_cell(0, 6, text)
    if user_password is not None or owner_password is not None:
        pdf.set_encryption(
            owner_password=owner_password or "owner-secret", user_password=user_password or ""
        )
    return bytes(pdf.output())


def make_layout_pdf(
    pages: list[list[tuple[float, str]]], running_header: str | None = None
) -> bytes:
    """One page per list of (font size, paragraph) — larger sizes read as headings. An optional
    running header is printed at the top of every page, like a book's page header."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=False)
    for paragraphs in pages:
        pdf.add_page()
        if running_header:
            pdf.set_font("Helvetica", size=8)
            pdf.set_xy(10, 5)
            pdf.cell(0, 4, running_header)
            pdf.set_xy(10, 25)
        for size, text in paragraphs:
            pdf.set_font("Helvetica", size=size)
            pdf.multi_cell(0, size * 0.5, text)
            pdf.ln(size * 0.6)  # paragraph spacing, well above the line spacing
    return bytes(pdf.output())


def make_blank_pdf(pages: int = 1) -> bytes:
    """Pages with no text layer — what a scanned document looks like to text extraction."""
    pdf = FPDF()
    for _ in range(pages):
        pdf.add_page()
    return bytes(pdf.output())


def make_docx(items: list[tuple[str, str]]) -> bytes:
    """items: ("heading", text) | ("para", text) | ("table", "a|b;c|d")."""
    document = new_docx()
    for kind, value in items:
        if kind == "heading":
            document.add_heading(value, level=1)
        elif kind == "para":
            document.add_paragraph(value)
        elif kind == "table":
            rows = [row.split("|") for row in value.split(";")]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx(slides: list[tuple[str, list[str]]]) -> bytes:
    """One slide per (title, bullet lines), using the "Title and Content" layout."""
    presentation = Presentation()
    layout = presentation.slide_layouts[1]
    for title, bullets in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for bullet in bullets[1:]:
            body.add_paragraph().text = bullet
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_png() -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (8, 8), "white").save(buffer, format="PNG")
    return buffer.getvalue()
