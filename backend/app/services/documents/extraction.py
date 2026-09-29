"""EXTRACT step (docs/PROJECT_SPEC.md §17): file bytes → paragraph-level blocks that remember
where they came from (page/slide, nearest heading), so source identity is never lost.

Every failure is raised as ExtractionError with a message that is safe to show the user and
contains none of the document's content.
"""

import re
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from docx import Document as open_docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PasswordType, PdfReader

from app.models.enums import DocumentKind
from app.services.documents import pdf_layout
from app.services.documents.blocks import Block, Extraction

MAX_PDF_PAGES = 3000
# DOCX/PPTX are ZIP archives; these bound what a small upload may expand to (zip bombs).
MAX_UNCOMPRESSED_BYTES = 500 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 20_000


class ExtractionError(Exception):
    """str(error) is shown to the user as the reason a document FAILED."""


__all__ = ["Block", "Extraction", "ExtractionError", "extract"]


def extract(kind: DocumentKind, data: bytes) -> Extraction:
    if kind is DocumentKind.PDF:
        return _extract_pdf(data)
    if kind is DocumentKind.TEXT:
        return Extraction(_plain_blocks(_decode_text(data)), None, None)
    if kind is DocumentKind.MARKDOWN:
        return Extraction(_markdown_blocks(_decode_text(data)), None, None)
    if kind is DocumentKind.DOCX:
        return _extract_docx(data)
    if kind is DocumentKind.PPTX:
        return _extract_pptx(data)
    raise ExtractionError(
        "Text extraction from images and scans isn't supported yet. The file is kept in the "
        "course's repository."
    )


# --- shared helpers ---


_BLANK_LINE = re.compile(r"\n\s*\n")


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in _BLANK_LINE.split(text) if p.strip()]


def _plain_blocks(text: str) -> list[Block]:
    return [Block(p) for p in _paragraphs(text)]


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")  # never fails: every byte is a latin-1 character


def _as_utc(value: Any) -> datetime | None:
    if not isinstance(value, datetime):
        return None
    # File metadata dates without a zone are taken as UTC rather than guessed.
    return value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)


def _check_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise ExtractionError("This file is damaged and couldn't be opened.") from exc
    uncompressed = sum(entry.file_size for entry in entries)
    if len(entries) > MAX_ARCHIVE_ENTRIES or uncompressed > MAX_UNCOMPRESSED_BYTES:
        raise ExtractionError("This file is too large to process once uncompressed.")


# --- PDF ---


def _extract_pdf(data: bytes) -> Extraction:
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted and reader.decrypt("") == PasswordType.NOT_DECRYPTED:
            raise ExtractionError("This PDF is password-protected. Remove the password and retry.")
        page_count = len(reader.pages)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("This PDF couldn't be read. It may be damaged.") from exc
    if page_count > MAX_PDF_PAGES:
        raise ExtractionError(f"PDFs are limited to {MAX_PDF_PAGES} pages.")

    def plain_paragraphs(page_number: int) -> list[str]:
        try:
            return _paragraphs(reader.pages[page_number - 1].extract_text() or "")
        except Exception:
            return []  # one unreadable page shouldn't sink the whole document

    # Layout first (headings from font sizes, real paragraphs); plain text for any page, or the
    # whole file, the layout reader can't handle.
    pages = pdf_layout.read_pages(data)
    if pages is not None:
        blocks = pdf_layout.build_blocks(pages, plain_paragraphs)
    else:
        blocks = [
            Block(text, page=number)
            for number in range(1, page_count + 1)
            for text in plain_paragraphs(number)
        ]

    # No headings recognizable from the layout: fall back to the PDF's bookmarks, if any.
    if not any(block.is_heading for block in blocks):
        sections = _pdf_sections(reader)
        if sections:
            blocks = [
                replace(block, heading=_section_for_page(sections, (block.page or 1) - 1))
                for block in blocks
            ]

    try:
        created = _as_utc(reader.metadata.creation_date) if reader.metadata else None
    except Exception:
        created = None  # malformed metadata dates are common and not worth failing over
    return Extraction(blocks, page_count, created)


def _pdf_sections(reader: PdfReader) -> list[tuple[int, str]]:
    """Bookmarks as (0-based start page, title), in page order — the PDF's section structure."""
    entries: list[tuple[int, str]] = []

    def walk(items: Any) -> None:
        for item in items:
            if isinstance(item, list):
                walk(item)
                continue
            try:
                page_index = reader.get_destination_page_number(item)
                title = str(item.title).strip()
            except Exception:  # noqa: S112 - skip one malformed bookmark, keep the rest
                continue
            if title and page_index is not None and page_index >= 0:
                entries.append((page_index, title[:500]))

    try:
        walk(reader.outline)
    except Exception:
        return []
    entries.sort(key=lambda entry: entry[0])  # stable: same-page entries keep outline order
    return entries


def _section_for_page(sections: list[tuple[int, str]], page_index: int) -> str | None:
    current = None
    for start, title in sections:
        if start > page_index:
            break
        current = title
    return current


# --- Markdown ---

_MD_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$")


def _markdown_blocks(text: str) -> list[Block]:
    blocks: list[Block] = []
    heading: str | None = None
    buffer: list[str] = []
    in_code = False

    def flush() -> None:
        paragraph = "\n".join(buffer).strip()
        if paragraph:
            blocks.append(Block(paragraph, heading=heading))
        buffer.clear()

    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            in_code = not in_code
            buffer.append(line)
            continue
        match = None if in_code else _MD_HEADING.match(line)
        if match:
            flush()
            heading = match.group(1).strip()[:500]
            blocks.append(Block(heading, heading=heading, is_heading=True))
        elif not line.strip() and not in_code:
            flush()
        else:
            buffer.append(line)
    flush()
    return blocks


# --- DOCX ---


def _extract_docx(data: bytes) -> Extraction:
    _check_archive(data)
    try:
        document = open_docx(BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This Word document couldn't be read.") from exc

    blocks: list[Block] = []
    heading: str | None = None
    # Walk the body in document order, so tables stay between the paragraphs around them.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            # Built-in heading styles keep their English names in the file, whatever the
            # Word UI language (e.g. Italian "Titolo 1" is stored as "Heading 1").
            style = paragraph.style.name if paragraph.style is not None else ""
            if style and (style.startswith("Heading") or style == "Title"):
                heading = text[:500]
                blocks.append(Block(text, heading=heading, is_heading=True))
            else:
                blocks.append(Block(text, heading=heading))
        elif tag == "tbl":
            for row in Table(child, document).rows:
                line = _table_row_text(cell.text for cell in row.cells)
                if line:
                    blocks.append(Block(line, heading=heading))

    created = _as_utc(document.core_properties.created)
    return Extraction(blocks, None, created)


def _table_row_text(cells: Iterable[str]) -> str:
    values: list[str] = []
    for cell in cells:
        value = cell.strip()
        # Merged cells repeat their text in each covered cell; keep one copy.
        if value and (not values or values[-1] != value):
            values.append(value)
    return " | ".join(values)


# --- PPTX ---


def _extract_pptx(data: bytes) -> Extraction:
    _check_archive(data)
    try:
        presentation = Presentation(BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This PowerPoint file couldn't be read.") from exc

    blocks: list[Block] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        heading = title[:500] or None
        if heading:
            blocks.append(Block(heading, page=number, heading=heading, is_heading=True))
        for shape in _iter_shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        blocks.append(Block(text, page=number, heading=heading))
            elif shape.has_table:
                for row in shape.table.rows:
                    line = _table_row_text(cell.text for cell in row.cells)
                    if line:
                        blocks.append(Block(line, page=number, heading=heading))

    created = _as_utc(presentation.core_properties.created)
    return Extraction(blocks, len(presentation.slides), created)


def _iter_shapes(shapes: Any) -> Iterator[Any]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape
